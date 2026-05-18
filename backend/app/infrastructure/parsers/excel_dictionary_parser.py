"""Excel data dictionary parser. Converts .xlsx field definitions into FileSchemaDefinition."""

from __future__ import annotations

import math
from io import BytesIO
from pathlib import Path

import structlog
from openpyxl import load_workbook

from app.domain.synthetic.file_schema import FileFieldDefinition, FileSchemaDefinition
from app.domain.synthetic.value_objects import CompType, EncodingType, FileFormat

logger = structlog.get_logger()

# Maximum input limits
MAX_DICTIONARY_BYTES = 10_485_760  # 10 MB
MAX_DICTIONARY_ROWS = 10_000

# Column alias mappings (lowercase)
COLUMN_ALIASES: dict[str, list[str]] = {
    "field_name": ["field_name", "name", "field", "column_name", "column", "fieldname", "field name", "column name"],
    "data_type": ["data_type", "type", "datatype", "data type"],
    "length": ["length", "len", "size", "field_length", "field length"],
    "start_position": ["start_position", "start", "position", "offset", "start position"],
    "decimal_places": ["decimal_places", "decimals", "scale", "decimal places"],
    "nullable": ["nullable", "null", "is_nullable", "is nullable"],
    "pii_type": ["pii_type", "pii", "sensitive_type", "pii type", "sensitive type"],
    "fk_reference": ["fk_reference", "fk", "foreign_key", "fk reference", "foreign key"],
    "comp_type": ["comp_type", "comp", "usage", "computational", "comp type"],
}

# Data type normalization
TYPE_MAP: dict[str, str] = {
    "char": "alphanumeric",
    "varchar": "alphanumeric",
    "string": "alphanumeric",
    "alpha": "alphanumeric",
    "alphanumeric": "alphanumeric",
    "text": "alphanumeric",
    "num": "numeric",
    "number": "numeric",
    "int": "numeric",
    "integer": "numeric",
    "numeric": "numeric",
    "dec": "decimal",
    "decimal": "decimal",
    "float": "decimal",
    "pack": "packed_decimal",
    "packed": "packed_decimal",
    "packed_decimal": "packed_decimal",
    "bin": "binary",
    "binary": "binary",
}

# Comp type normalization
COMP_MAP: dict[str, str] = {
    "comp": CompType.COMP.value,
    "comp-0": CompType.COMP.value,
    "binary": CompType.COMP.value,
    "computational": CompType.COMP.value,
    "comp-3": CompType.COMP_3.value,
    "packed": CompType.COMP_3.value,
    "packed-decimal": CompType.COMP_3.value,
    "computational-3": CompType.COMP_3.value,
}


class ExcelDictionaryParser:
    """Parse an Excel data dictionary (.xlsx) into a FileSchemaDefinition."""

    def parse(
        self,
        file_bytes: bytes,
        name: str = "dictionary",
        target_format: FileFormat = FileFormat.CSV,
        target_encoding: EncodingType = EncodingType.UTF8,
    ) -> FileSchemaDefinition:
        """Parse Excel bytes into a FileSchemaDefinition.

        Args:
            file_bytes: Raw .xlsx file content.
            name: Name for the resulting schema.
            target_format: Target file format.
            target_encoding: Target encoding.

        Returns:
            FileSchemaDefinition with parsed fields.

        Raises:
            ValueError: If input exceeds limits or required columns are missing.
        """
        if len(file_bytes) > MAX_DICTIONARY_BYTES:
            raise ValueError(f"Excel file exceeds size limit (max {MAX_DICTIONARY_BYTES // 1_048_576}MB)")

        wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
        ws = wb.active
        if ws is None:
            raise ValueError("Excel workbook has no active sheet")

        # Find header row (scan first 5 rows)
        column_map = self._find_header_row(ws)

        # Validate required columns
        missing = []
        for req in ("field_name", "data_type", "length"):
            if req not in column_map:
                missing.append(req)
        if missing:
            raise ValueError(f"Required columns not found: {', '.join(missing)}")

        # Parse data rows
        fields: list[FileFieldDefinition] = []
        current_position = 0
        row_count = 0

        for row in ws.iter_rows(min_row=column_map["_header_row"] + 1, values_only=False):
            row_count += 1
            if row_count > MAX_DICTIONARY_ROWS:
                raise ValueError(f"Excel dictionary exceeds row limit (max {MAX_DICTIONARY_ROWS})")

            row_values = {i: cell.value for i, cell in enumerate(row)}

            # Get field name — skip empty rows
            field_name = self._get_cell(row_values, column_map, "field_name")
            if not field_name:
                continue

            field_name = str(field_name).strip()
            if not field_name:
                continue

            # Get required fields
            raw_type = str(self._get_cell(row_values, column_map, "data_type") or "").strip().lower()
            raw_length = self._get_cell(row_values, column_map, "length")

            if not raw_type:
                raise ValueError(f"Missing data_type for field '{field_name}'")

            try:
                length = int(raw_length)
            except (TypeError, ValueError):
                raise ValueError(f"Invalid length '{raw_length}' for field '{field_name}'")

            # Normalize data type
            data_type = TYPE_MAP.get(raw_type)
            if not data_type:
                raise ValueError(
                    f"Unknown data type '{raw_type}' for field '{field_name}'. "
                    f"Valid types: {', '.join(sorted(set(TYPE_MAP.values())))}"
                )

            # Optional fields
            raw_start = self._get_cell(row_values, column_map, "start_position")
            raw_decimals = self._get_cell(row_values, column_map, "decimal_places")
            raw_nullable = self._get_cell(row_values, column_map, "nullable")
            raw_pii = self._get_cell(row_values, column_map, "pii_type")
            raw_fk = self._get_cell(row_values, column_map, "fk_reference")
            raw_comp = self._get_cell(row_values, column_map, "comp_type")

            # Parse comp_type
            comp_type_str = CompType.NONE.value
            if raw_comp:
                comp_key = str(raw_comp).strip().lower().replace("_", "-")
                comp_type_str = COMP_MAP.get(comp_key, CompType.NONE.value)

            # Parse decimal_places
            decimal_places = 0
            if raw_decimals is not None:
                try:
                    decimal_places = int(raw_decimals)
                except (TypeError, ValueError):
                    decimal_places = 0

            # Compute byte_length
            byte_length = self._compute_byte_length(data_type, length, decimal_places, comp_type_str)

            # Parse start_position
            if raw_start is not None:
                try:
                    start_position = int(raw_start)
                except (TypeError, ValueError):
                    start_position = current_position
            else:
                start_position = current_position

            # Parse nullable
            nullable = True
            if raw_nullable is not None:
                nullable_str = str(raw_nullable).strip().lower()
                nullable = nullable_str not in ("false", "no", "n", "0")

            # Parse pii_type
            pii_type = str(raw_pii).strip() if raw_pii else None

            # Parse fk_reference
            fk_reference = None
            if raw_fk:
                fk_str = str(raw_fk).strip()
                if "." in fk_str:
                    parts = fk_str.split(".", 1)
                    fk_reference = (parts[0], parts[1])

            field_def = FileFieldDefinition(
                name=field_name,
                data_type=data_type,
                length=length,
                byte_length=byte_length,
                start_position=start_position,
                decimal_places=decimal_places,
                comp_type=comp_type_str,
                nullable=nullable,
                pii_type=pii_type,
                fk_reference=fk_reference,
            )
            fields.append(field_def)

            # Advance auto-position
            current_position = start_position + byte_length

        wb.close()

        if not fields:
            raise ValueError("No valid field definitions found in Excel dictionary")

        record_length = sum(f.byte_length for f in fields if f.redefines is None)

        schema = FileSchemaDefinition(
            name=name,
            fields=fields,
            file_format=target_format.value,
            encoding=target_encoding.value,
            record_length=record_length,
            metadata={"source_type": "excel_dictionary"},
        )

        logger.info(
            "excel_dictionary_parsed",
            name=name,
            field_count=len(fields),
            record_length=record_length,
        )

        return schema

    def parse_file(self, file_path: Path, **kwargs) -> FileSchemaDefinition:
        """Parse a dictionary file from disk."""
        data = file_path.read_bytes()
        if "name" not in kwargs:
            kwargs["name"] = file_path.stem
        return self.parse(data, **kwargs)

    # ── Internal Methods ────────────────────────────────────────────

    def _find_header_row(self, ws) -> dict[str, int]:
        """Scan first 5 rows for a header row. Returns {canonical_name: column_index, "_header_row": row_number}."""
        for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=False), start=1):
            cell_values = []
            for cell in row:
                val = str(cell.value).strip().lower().replace(" ", "_") if cell.value else ""
                cell_values.append(val)

            # Check if this row contains at least field_name
            mapping: dict[str, int] = {}
            for col_idx, cell_val in enumerate(cell_values):
                if not cell_val:
                    continue
                # Normalize: strip underscores/spaces and compare
                normalized = cell_val.replace("_", " ").strip()
                for canonical, aliases in COLUMN_ALIASES.items():
                    if normalized in aliases or cell_val in [a.replace(" ", "_") for a in aliases]:
                        mapping[canonical] = col_idx
                        break

            if "field_name" in mapping:
                mapping["_header_row"] = row_idx
                return mapping

        raise ValueError(
            "Could not find header row in first 5 rows. "
            "Expected a row containing 'field_name' (or alias: name, field, column_name, column)"
        )

    def _get_cell(self, row_values: dict[int, any], column_map: dict[str, int], column_name: str):
        """Get a cell value from a row using the column mapping."""
        col_idx = column_map.get(column_name)
        if col_idx is None:
            return None
        return row_values.get(col_idx)

    def _compute_byte_length(
        self, data_type: str, length: int, decimal_places: int, comp_type: str
    ) -> int:
        """Compute physical byte length — same logic as CopybookParser."""
        total_digits = length

        if comp_type == CompType.COMP_3.value:
            nibbles = total_digits + 1
            return math.ceil(nibbles / 2)
        elif comp_type == CompType.COMP.value:
            if total_digits <= 4:
                return 2
            elif total_digits <= 9:
                return 4
            else:
                return 8
        else:
            return length
